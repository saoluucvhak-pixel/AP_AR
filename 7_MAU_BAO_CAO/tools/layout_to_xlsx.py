#!/usr/bin/env python3
"""Dựng tệp Excel XEM TRƯỚC từ cùng bố cục (reportLayout_) mà Backend áp lên Google Sheets.
Đây là bản mô phỏng bằng openpyxl để xem thiết kế; tệp thật do Google Sheets xuất ra có thể khác chút về phông/độ rộng."""
import json,sys,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as L
R=dict(brand='1F3A5F',headFill='1F3A5F',headInk='FFFFFF',zebra='F6F8FB',group='E3EAF4',subtotal='EEF2F8',total='D5E0F0',border='B9C3D3',muted='6B7280',danger='B00020')
def fill(c): return PatternFill('solid',fgColor=c)
thin=Side(style='thin',color=R['border']); box=Border(left=thin,right=thin,top=thin,bottom=thin)
def build(lay,out,sheet,help_):
    wb=Workbook(); ws=wb.active; ws.title=sheet; nc=lay['nc']; ws.sheet_view.showGridLines=False
    for i,w in enumerate(lay['widths']): ws.column_dimensions[L(i+1)].width=max(6,w/7.0)
    for ri,r in enumerate(lay['rows'],1):
        k=r['kind']; fm=r.get('fmts') or ['@']*nc
        for ci,v in enumerate(r['cells'],1):
            c=ws.cell(row=ri,column=ci,value=(None if v=='' else v)); c.font=Font(name='Arial',size=10); c.alignment=Alignment(vertical='center')
            if k in('data','group','subtotal','total') and isinstance(v,(int,float)): c.number_format=fm[ci-1].replace('0;[Red]-#,##0;0','0;[Red]-#,##0;0')
        if r.get('height'): ws.row_dimensions[ri].height=r['height']*0.75
        rng=[ws.cell(row=ri,column=c) for c in range(1,nc+1)]
        def st(**kw):
            for c in rng:
                for a,b in kw.items():
                    if a=='font': c.font=b
                    elif a=='fill': c.fill=b
                    elif a=='al': c.alignment=b
        if k=='mast1':
            p=r['parts']
            for i,(c0,sp) in enumerate(p):
                cell=ws.cell(row=ri,column=c0+1)
                if i==0: cell.font=Font(name='Arial',size=11,bold=True,color=R['brand']); cell.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
                elif i==1: cell.font=Font(name='Arial',size=r.get('titleSize',15),bold=True,color=R['brand']); cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                else: cell.font=Font(name='Arial',size=8,bold=True,color=R['danger']); cell.alignment=Alignment(horizontal='right',vertical='center',wrap_text=True)
        elif k=='mast2':
            p=r['parts']
            for c in rng:
                c.font=Font(name='Arial',size=8,color=R['muted']); c.border=Border(bottom=Side(style='thin',color=R['brand']))
            ws.cell(row=ri,column=p[0][0]+1).alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
            ws.cell(row=ri,column=p[1][0]+1).alignment=Alignment(horizontal='right',vertical='center')
        elif k=='band': st(font=Font(name='Arial',size=9,italic=True,color='374151'),fill=fill('FAFBFD'),al=Alignment(wrap_text=True,vertical='center'))
        elif k=='company':
            st(font=Font(name='Arial',size=13,bold=True,color=R['brand']))
            left=r['parts'][1][0] if len(r['parts'])>1 else nc
            if left<nc: ws.cell(row=ri,column=left+1).font=Font(name='Arial',size=9,color=R['muted']); ws.cell(row=ri,column=left+1).alignment=Alignment(horizontal='right',vertical='center')
        elif k=='contact': st(font=Font(name='Arial',size=9,color=R['muted']))
        elif k=='title': st(font=Font(name='Arial',size=16,bold=True,color=R['brand']),al=Alignment(horizontal='center',vertical='center'))
        elif k=='subtitle': st(font=Font(name='Arial',size=10,color='374151'),al=Alignment(horizontal='center',vertical='center',wrap_text=True))
        elif k=='status': st(font=Font(name='Arial',size=9,bold=True,color=R['danger']),al=Alignment(horizontal='center',vertical='center'))
        elif k=='summaryHead': st(font=Font(name='Arial',size=10,bold=True,color=R['brand']),fill=fill(R['group']))
        elif k=='summary': st(fill=fill('FAFBFD'),al=Alignment(wrap_text=True,vertical='center'))
        elif k=='header': st(font=Font(name='Arial',size=10,bold=True,color=R['headInk']),fill=fill(R['headFill']),al=Alignment(horizontal='center',vertical='center',wrap_text=True))
        elif k=='group': st(font=Font(name='Arial',size=10,bold=True),fill=fill(R['group']))
        elif k=='subtotal': st(font=Font(name='Arial',size=10,bold=True),fill=fill(R['subtotal']))
        elif k=='total': st(font=Font(name='Arial',size=10,bold=True),fill=fill(R['total']))
        elif k=='data' and r.get('zebra'): st(fill=fill(R['zebra']))
        elif k=='words': st(font=Font(name='Arial',size=10,bold=True,italic=True),al=Alignment(wrap_text=True,vertical='center'))
        elif k=='note': st(font=Font(name='Arial',size=9,italic=True,color='374151'),al=Alignment(wrap_text=True,vertical='center'))
        elif k in('sig','sigName'): st(font=Font(name='Arial',size=10,bold=True))
        elif k=='sigHint': st(font=Font(name='Arial',size=9,italic=True,color=R['muted']))
        elif k=='footer': st(font=Font(name='Arial',size=8,color=R['muted']),al=Alignment(horizontal='center',vertical='center'))
    for m in lay['merges']:
        r0,c0,nr,ncc=m
        if nr*ncc>1: ws.merge_cells(start_row=r0,start_column=c0,end_row=r0+nr-1,end_column=c0+ncc-1)
    h,lo=lay['headerRow'],lay['lastData']
    for ri in range(h,lo+1):
        for ci in range(1,nc+1): ws.cell(row=ri,column=ci).border=box
        kind=lay['rows'][ri-1]['kind']
        if kind=='total':
            for ci in range(1,nc+1): ws.cell(row=ri,column=ci).border=Border(left=thin,right=thin,top=Side(style='medium',color=R['brand']),bottom=Side(style='medium',color=R['brand']))
        if kind in('data','group','subtotal','total') and ri>h:
            pass
    return wb,ws
def finish(wb,ws,lay,hp):
    nc=lay['nc']; ws.freeze_panes=('A%d'%(lay['frozenRows']+1)) if lay['frozenRows'] else None
    # căn lề theo kiểu cột: số bên phải
    for ri in range(lay['firstData'],lay['lastData']+1):
        for ci in range(1,nc+1):
            c=ws.cell(row=ri,column=ci)
            if isinstance(c.value,(int,float)): c.alignment=Alignment(horizontal='right',vertical='top')
            else: c.alignment=Alignment(vertical='top',wrap_text=True)
    ws.sheet_properties.tabColor=R['brand']
    ws.page_setup.orientation='landscape' if lay['landscape'] else 'portrait'; ws.page_setup.paperSize=9
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.print_title_rows='1:%d'%lay['headerRow'] if lay['frozenRows'] else '%d:%d'%(lay['headerRow'],lay['headerRow'])
    ws.print_options.horizontalCentered=True
    ws.oddFooter.center.text='Trang &P / &N'; ws.oddFooter.center.size=8
    ws.page_margins.left=ws.page_margins.right=0.4; ws.page_margins.top=0.5; ws.page_margins.bottom=0.6
    hs=wb.create_sheet('Hướng dẫn đọc'); hs.sheet_view.showGridLines=False
    for i,(a,b) in enumerate(hp,1):
        hs.cell(row=i,column=1,value=a).font=Font(name='Arial',bold=True,size=13 if i==1 else 10,color=R['brand'] if i==1 else '000000')
        c=hs.cell(row=i,column=2,value=b); c.font=Font(name='Arial',size=10); c.alignment=Alignment(wrap_text=True,vertical='top'); hs.cell(row=i,column=1).alignment=Alignment(vertical='top',wrap_text=True)
    hs.column_dimensions['A'].width=24; hs.column_dimensions['B'].width=95
if __name__=='__main__':
    d=json.load(open(sys.argv[1])); outdir=sys.argv[2]
    for code,lay in d.items():
        wb,ws=build(lay,None,code,lay['help']); finish(wb,ws,lay,lay['help']); wb.save('%s/%s_mau_xem_truoc.xlsx'%(outdir,code)); print(code)
